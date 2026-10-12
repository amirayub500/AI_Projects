"""
"Download as Excel" for "/" answers.

The page sends back the tables it received with the answer (the data behind
each chart and each database query). This file turns them into an .xlsx:
one sheet per table, plus an "About" sheet with the question, date and SQL.
"""

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_TABLES = 20
MAX_ROWS = 5000          # per table
MAX_COLUMNS = 50

HEADER_FILL = PatternFill("solid", start_color="0B1530")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _safe_cell(value):
    """Numbers stay numbers. Text that starts like a formula is made plain text,
    so a cell can never run as an Excel formula."""
    if value is None or isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value)
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


def _sheet_name(title, used):
    """Excel sheet names: max 31 characters, no []:*?/\\ and unique."""
    name = re.sub(r"[\[\]:*?/\\]", " ", str(title or "Data")).strip()[:28] or "Data"
    base, n = name, 2
    while name.lower() in used:
        name = f"{base[:25]} ({n})"
        n += 1
    used.add(name.lower())
    return name


def build_workbook(question, tables):
    """Return the .xlsx file as bytes."""
    wb = Workbook()

    about = wb.active
    about.title = "About"
    about.append(["AK Data Solutions - Research AI Agent export"])
    about["A1"].font = Font(bold=True, size=13)
    about.append(["Question", _safe_cell(question)])
    about.append(["Exported", datetime.now().strftime("%Y-%m-%d %H:%M")])
    about.append([])
    about.append(["Sheet", "Rows", "Source query (SQL)"])
    for cell in about[5]:
        cell.font = Font(bold=True)

    used = {"about"}
    for table in tables[:MAX_TABLES]:
        columns = [str(c) for c in (table.get("columns") or [])][:MAX_COLUMNS]
        rows = (table.get("rows") or [])[:MAX_ROWS]
        if not columns:
            continue

        ws = wb.create_sheet(_sheet_name(table.get("title"), used))
        ws.append(columns)
        for cell in ws[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        for row in rows:
            ws.append([_safe_cell(v) for v in list(row)[:len(columns)]])

        ws.freeze_panes = "A2"
        for i, col in enumerate(columns, start=1):
            width = max([len(col)] + [len(str(r[i - 1])) for r in rows[:200] if len(r) >= i and r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
            if any(isinstance(r[i - 1], float) for r in rows[:50] if len(r) >= i):
                for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                    for c in cell:
                        c.number_format = "#,##0.00"

        about.append([ws.title, len(rows), _safe_cell(table.get("sql", ""))])

    about.column_dimensions["A"].width = 30
    about.column_dimensions["B"].width = 12
    about.column_dimensions["C"].width = 90
    for row in about.iter_rows(min_row=6):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
