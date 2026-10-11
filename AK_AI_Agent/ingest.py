"""
ingest.py  -  convert the client's files to text and load them into the database
================================================================================
Run it on a folder; it walks all sub-folders.

    python ingest.py data/inbox                 import everything new
    python ingest.py data/inbox --reimport      import again even if already loaded
    python ingest.py --status                   show what is loaded

In Docker on the server (files placed in ~/AI_Projects/AK_AI_Agent/data/inbox):
    docker compose run --rm app python ingest.py data/inbox

Supported files:
    .pdf                 text per page; scanned pages are OCR'd if Tesseract is installed
    .docx                paragraphs + tables (stored as "part 1", "part 2", ...)
    .txt .md             plain text
    .csv .xlsx .xlsm     each sheet becomes a SQL table (see company_db.add_table)

Every file gets an ID = hash of its contents, so the same file is never imported
twice, even if it is renamed or copied to another folder.
"""

import argparse
import hashlib
import os
import sys
import time
from contextlib import closing

import company_db


DOC_TYPES = {".pdf", ".docx", ".txt", ".md"}
TABLE_TYPES = {".csv", ".xlsx", ".xlsm", ".xls"}

# A PDF page with fewer characters than this is treated as a scan -> OCR
MIN_PAGE_CHARS = 25

# Characters per "part" for files without pages (Word, text)
PART_SIZE = 3000


# =====================================================================
# Helpers
# =====================================================================

def doc_id_for(path):
    """Stable ID from the file's bytes: same content -> same ID."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()[:16]


def ocr_available():
    """True if pytesseract and the Tesseract program are both installed."""
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def split_parts(text):
    """Split long text without pages into labelled parts of about PART_SIZE chars."""
    parts, current = [], ""

    for paragraph in text.split("\n"):
        if len(current) + len(paragraph) > PART_SIZE and current:
            parts.append(current)
            current = ""
        current += paragraph + "\n"

    if current.strip():
        parts.append(current)

    return [(f"part {i}", p) for i, p in enumerate(parts, start=1)]


# =====================================================================
# Converters: file -> list of (label, text), plus status and note
# =====================================================================

def read_pdf(path, use_ocr):
    try:
        import pymupdf as fitz    # PyMuPDF (newer name)
    except ImportError:
        import fitz               # PyMuPDF (older name)

    pages, ocr_pages, empty_pages = [], 0, 0

    with fitz.open(path) as pdf:
        for number, page in enumerate(pdf, start=1):
            text = page.get_text("text")

            if len(text.strip()) < MIN_PAGE_CHARS:
                if use_ocr:
                    text = ocr_page(page)
                    ocr_pages += 1
                else:
                    empty_pages += 1

            pages.append((f"page {number}", text))

    total = len(pages)

    if total and empty_pages == total:
        return pages, "needs_ocr", "scanned PDF; install Tesseract and re-import"

    note = []
    if ocr_pages:
        note.append(f"{ocr_pages} page(s) OCR'd")
    if empty_pages:
        note.append(f"{empty_pages} page(s) without text")

    return pages, "ok", "; ".join(note)


def ocr_page(page):
    """Render a PDF page as an image and read it with Tesseract."""
    import io
    import pytesseract
    from PIL import Image

    pixmap = page.get_pixmap(dpi=200)
    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
    return pytesseract.image_to_string(image)


def read_docx(path):
    import docx   # python-docx

    document = docx.Document(path)
    lines = [p.text for p in document.paragraphs if p.text.strip()]

    # Tables inside the Word file: one line per row, cells separated by " | "
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    return split_parts("\n".join(lines)), "ok", ""


def read_text(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return split_parts(f.read()), "ok", ""


def read_tables(path):
    """Spreadsheet -> list of (sheet_name or None, DataFrame)."""
    import pandas as pd

    if path.lower().endswith(".csv"):
        try:
            return [(None, pd.read_csv(path))]
        except UnicodeDecodeError:
            return [(None, pd.read_csv(path, encoding="latin-1"))]

    sheets = pd.read_excel(path, sheet_name=None)
    single = len(sheets) == 1
    return [(None if single else name, df) for name, df in sheets.items()]


# =====================================================================
# Import one file
# =====================================================================

def ingest_file(conn, path, reimport, use_ocr):
    """Import one file. Returns a short result string for the progress log."""
    file_name = os.path.basename(path)
    ext = os.path.splitext(file_name)[1].lower()
    doc_id = doc_id_for(path)

    if company_db.document_exists(conn, doc_id):
        if not reimport:
            return "skipped (already imported)"
        company_db.delete_document(conn, doc_id)

    try:
        if ext in TABLE_TYPES:
            names = []
            for sheet, df in read_tables(path):
                names.append(company_db.add_table(conn, df, file_name, sheet))
            company_db.add_document(
                conn, doc_id, file_name, path, ext[1:], [], "ok",
                "tables: " + ", ".join(names),
            )
            return "table(s) " + ", ".join(names)

        if ext == ".pdf":
            pages, status, note = read_pdf(path, use_ocr)
        elif ext == ".docx":
            pages, status, note = read_docx(path)
        else:
            pages, status, note = read_text(path)

        company_db.add_document(conn, doc_id, file_name, path, ext[1:], pages, status, note)
        return f"{status}, {len(pages)} page(s)/part(s)" + (f" ({note})" if note else "")

    except Exception as exc:
        # Record the failure and carry on with the next file
        company_db.add_document(conn, doc_id, file_name, path, ext[1:], [], "error", str(exc)[:300])
        return f"ERROR: {exc}"


def find_files(targets):
    for target in targets:
        if os.path.isfile(target):
            yield target
            continue

        for root, _, files in os.walk(target):
            for name in sorted(files):
                if name.startswith("~$"):      # Office temp/lock files
                    continue
                if os.path.splitext(name)[1].lower() in DOC_TYPES | TABLE_TYPES:
                    yield os.path.join(root, name)


def print_status(path):
    if not os.path.exists(path):
        print(f"No database yet at {path}")
        return

    with closing(company_db.connect_write(path)) as conn:
        print(f"Database: {path}\n")

        for status, count in conn.execute(
            "SELECT status, COUNT(*) FROM documents GROUP BY status ORDER BY status"
        ):
            print(f"  {status:10} {count} file(s)")

        print("\nTables:")
        print(company_db.describe_tables(path))

        problems = conn.execute(
            "SELECT file_name, status, note FROM documents WHERE status != 'ok' LIMIT 20"
        ).fetchall()
        if problems:
            print("\nNeeds attention:")
            for name, status, note in problems:
                print(f"  [{status}] {name}: {note}")


def main():
    parser = argparse.ArgumentParser(description="Import client files into the company database.")
    parser.add_argument("targets", nargs="*", help="Folders or files to import")
    parser.add_argument("--reimport", action="store_true", help="Import files again even if already loaded")
    parser.add_argument("--no-ocr", action="store_true", help="Skip OCR for scanned PDFs")
    parser.add_argument("--status", action="store_true", help="Show what is loaded and stop")
    parser.add_argument("--db", default=company_db.DB_PATH, help="Database file path")
    args = parser.parse_args()

    if args.status:
        print_status(args.db)
        return

    if not args.targets:
        parser.error("give a folder or file to import, or --status")

    use_ocr = not args.no_ocr and ocr_available()
    print(f"Database: {args.db}")
    print(f"OCR for scanned PDFs: {'on' if use_ocr else 'off (Tesseract not found or --no-ocr)'}\n")

    started = time.monotonic()
    count = 0

    with closing(company_db.connect_write(args.db)) as conn:
        for path in find_files(args.targets):
            count += 1
            result = ingest_file(conn, path, args.reimport, use_ocr)
            conn.commit()     # save after every file, so a crash loses nothing
            print(f"[{count}] {path}: {result}")

    print(f"\nDone: {count} file(s) in {time.monotonic() - started:.1f} s")
    print_status(args.db)


if __name__ == "__main__":
    sys.exit(main())
